import openpyxl,json,sys,datetime
wb=openpyxl.load_workbook(sys.argv[1],data_only=True)
def d(v):
    if isinstance(v,datetime.datetime): return v.strftime('%Y-%m-%d')
    return str(v)[:10] if v else None
def n(v):
    try: return float(v)
    except: return 0.0
log=[]
for r in wb['Log'].iter_rows(min_row=3,values_only=True):
    if not r[0] or not isinstance(r[0],datetime.datetime): continue
    full=n(r[2]); split=(r[10]=='Yes'); g=int(n(r[11]) or 1)
    share=round(full/g,2) if split and g>0 else full
    log.append(dict(d=d(r[0]),ds=r[1] or '',a=full,c=r[3] or 'Uncategorised',t=r[4] or '',m=r[5] or '',rc=r[6]=='Yes',p=r[7] or '',b=r[8] or '',sp=split,g=g,s=share,rv=round(full-share,2) if split else 0,pp=r[14] or '',tr=r[15] or '',lp=r[16] or '',la=n(r[17]),rt=r[18] or '',nt=r[19] or '',st=r[22] or ''))
home=[]
for r in wb['Home Construction'].iter_rows(min_row=2,values_only=True):
    if not r[0]: continue
    home.append(dict(d=d(r[0]),h=r[1],ds=r[2],q=n(r[3]),rate=n(r[4]),u=r[5],a=n(r[8]),pd=n(r[9]),pe=n(r[10]),c=r[11],f=r[12],st=r[13]))
bud=[]
for r in wb['Budget vs Actual'].iter_rows(min_row=4,values_only=True):
    if r[0] and r[0]!='Total Fixed Spends' and isinstance(r[1],(int,float)): bud.append(dict(i=r[0],a=r[1],t=r[2]))
loans=[dict(p=r[0],ds=r[1],d=d(r[2]),a=n(r[3]),r=n(r[5]),st=r[7]) for r in wb['Loan Tracker'].iter_rows(min_row=3,values_only=True) if r[0] and r[0]!='TOTAL OUTSTANDING']
inv=[dict(d=d(r[0]),t=r[1],ds=r[2],a=n(r[3]),nt=r[6]) for r in wb['Investments'].iter_rows(min_row=3,values_only=True) if r[0] and r[1] and isinstance(r[3],(int,float))]
out=dict(log=log,home=home,budget=bud,loans=loans,inv=inv,asof=max(x['d'] for x in log))
json.dump(out,open(sys.argv[2],'w'),separators=(',',':'),ensure_ascii=False)
print(len(log),len(home),len(bud),len(loans),len(inv),out['asof'])
